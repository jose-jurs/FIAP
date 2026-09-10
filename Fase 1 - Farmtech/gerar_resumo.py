# -*- coding: utf-8 -*-
"""
Gera o resumo do artigo (Formacao Social) em .docx no formato exigido:
Arial 11, espacamento simples (1), margens direita/esquerda de 2 cm.
Artigo: "Uso de veiculos aereos nao tripulados (VANT) em Agricultura de
Precisao" - Lucio A. C. Jorge e Ricardo Y. Inamasu (Embrapa Instrumentacao).
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

TITULO = ("Resumo - Uso de Veiculos Aereos Nao Tripulados (VANT) em "
          "Agricultura de Precisao")

# Texto do resumo (redacao propria a partir da leitura integral do artigo).
PARAGRAFOS = [
    "O capitulo de Lucio Andre de Castro Jorge e Ricardo Y. Inamasu, "
    "pesquisadores da Embrapa Instrumentacao (Sao Carlos-SP), discute o uso "
    "de Veiculos Aereos Nao Tripulados (VANTs), tambem chamados de drones, "
    "como ferramenta de apoio a agricultura de precisao. Os autores destacam "
    "que o interesse por essas aeronaves cresce no mundo e, em especial, no "
    "Brasil, impulsionado pela reducao do custo e do tamanho dos "
    "equipamentos e pela necessidade de otimizar a producao. O objetivo do "
    "texto nao e esgotar o tema, mas apresentar o contexto de desenvolvimento "
    "dos VANTs no pais, descrever as iniciativas em curso e avaliar as "
    "questoes ligadas a integracao segura dessas aeronaves na atividade "
    "agricola.",

    "No panorama nacional, o trabalho relata que as pesquisas ainda sao "
    "recentes e concentram-se em aplicacoes civis como vigilancia, inspecao "
    "de linhas de transmissao, monitoramento de safras e controle de pragas e "
    "queimadas. A Embrapa investiu no desenvolvimento de plataformas e "
    "aeronaves capazes de operar em condicoes adversas de campo, e "
    "universidades deram origem a empresas nacionais (como XMobots, Airship e "
    "Skydrones) que passaram a atender o mercado. Um ponto central e a "
    "regulamentacao: cabe a ANAC autorizar o uso do espaco aereo, homologar "
    "os equipamentos e padronizar as formas de operacao, o que ainda "
    "representa uma barreira a difusao da tecnologia.",

    "Tecnicamente, os autores explicam que um VANT depende de um sistema de "
    "controle capaz de estabilizar a aeronave e conduzi-la por uma rota "
    "programada. Para isso, combinam-se GPS, unidade de navegacao inercial "
    "(IMU), magnetometro e um piloto automatico (AFCS) que recebe comandos da "
    "estacao de solo (GCS). O texto tambem classifica os tipos de aeronave e "
    "detalha os sensores empregados, baseados na espectroscopia de "
    "reflectancia, ou seja, na medida da radiacao eletromagnetica refletida "
    "pela vegetacao nas faixas do visivel, do infravermelho e em sensores "
    "hiperespectrais, alem da deteccao de fluorescencia, que permitem "
    "inferir o estado nutricional e hidrico das plantas.",

    "Uma parte relevante do capitulo trata da metodologia de uso em campo. Os "
    "autores mostram como a altitude de voo determina a resolucao da imagem "
    "(centimetros por pixel) e a area coberta, e discutem cuidados para "
    "evitar borramento, como o ajuste da velocidade do obturador, a fixacao "
    "do balanco de branco e o uso do formato RAW. Recomendam sobreposicao "
    "lateral minima de 40% (idealmente 60%) entre as fotos para garantir a "
    "qualidade dos mosaicos, e descrevem as etapas seguintes do fluxo de "
    "trabalho: planejamento da missao, captura das imagens, processamento e "
    "mosaicagem, analise em ferramenta SIG (GIS) e geracao de relatorios.",

    "Nas conclusoes, os autores reforcam que o uso de VANTs na agricultura de "
    "precisao tende a crescer, ja que a tecnologia se torna mais acessivel, "
    "confiavel e precisa, colocando ao alcance do produtor recursos de "
    "sensoriamento remoto antes restritos a satelites e avioes tripulados. "
    "Ressaltam, contudo, que a consolidacao dessa pratica depende do avanco "
    "da regulamentacao pela ANAC, com a homologacao das aeronaves e a "
    "obtencao das autorizacoes de voo. Do ponto de vista da formacao social, "
    "o artigo evidencia como a inovacao tecnologica no campo articula "
    "ganhos de produtividade e sustentabilidade com a necessidade de marcos "
    "legais e responsabilidade no uso do espaco aereo.",
]


def build():
    doc = Document()

    # Margens: direita e esquerda em 2 cm (topo/base padrao)
    for secao in doc.sections:
        secao.left_margin = Cm(2)
        secao.right_margin = Cm(2)

    # Estilo base: Arial 11, espacamento simples
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    # Titulo
    titulo = doc.add_paragraph()
    run = titulo.add_run(TITULO)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(6)

    # Linha de autoria/fonte
    fonte = doc.add_paragraph()
    r2 = fonte.add_run(
        "Artigo original: JORGE, L. A. C.; INAMASU, R. Y. Uso de veiculos "
        "aereos nao tripulados (VANT) em Agricultura de Precisao. In: "
        "Agricultura de Precisao: resultados de um novo olhar. Embrapa, "
        "cap. 8, p. 109-134."
    )
    r2.italic = True
    r2.font.name = "Arial"
    r2.font.size = Pt(9)
    fonte.paragraph_format.space_after = Pt(10)

    # Corpo justificado
    for texto in PARAGRAFOS:
        p = doc.add_paragraph(texto)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    doc.save("Resumo_Artigo_VANT.docx")
    print("[OK] Resumo_Artigo_VANT.docx gerado.")


if __name__ == "__main__":
    build()
